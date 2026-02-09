"""
Excel (.xlsx / .xlsm / .xltx / .xltm) to Markdown converter.

Handles merged cells, formulas, hyperlinks, and large workbooks.
Uses ``openpyxl`` for reading and ``tabulate`` for table formatting.

Ported from the original ``convert_xlsx_to_md.py``.
"""

import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from .base import BaseConverter, ConversionResult
from .registry import ConverterRegistry

logger = logging.getLogger(__name__)

# File size threshold for switching to read-only streaming mode (bytes).
_LARGE_FILE_THRESHOLD = 10 * 1024 * 1024  # 10 MB


@ConverterRegistry.register
class XlsxConverter(BaseConverter):
    """Convert modern Excel files to Markdown with merged-cell awareness."""

    EXTENSIONS = [".xlsx", ".xlsm", ".xltx", ".xltm"]
    DISPLAY_NAME = "Excel (XLSX)"
    REQUIRED_PACKAGES = {"openpyxl": "openpyxl", "tabulate": "tabulate"}
    MAX_MEMORY_WARNING_MB = 50

    def __init__(self, **options: Any) -> None:
        super().__init__(**options)
        self.preserve_formulas: bool = options.get("preserve_formulas", False)
        self.include_metadata: bool = options.get("include_metadata", True)
        self.max_column_width: int = options.get("max_column_width", 0)  # 0 = no limit
        self.empty_cell_placeholder: str = options.get("empty_cell_placeholder", "")
        self.include_hidden_sheets: bool = options.get("include_hidden_sheets", False)
        self.sheets: Optional[List[str]] = options.get("sheets")

    # -- public API -----------------------------------------------------------

    def convert_file(
        self,
        input_path: Path,
        output_path: Optional[Path] = None,
    ) -> ConversionResult:
        from openpyxl import load_workbook

        input_path = Path(input_path)
        output_path = self.resolve_output_path(input_path, output_path)

        # Decide whether to use read-only mode for large files
        file_size = input_path.stat().st_size
        use_readonly = file_size > _LARGE_FILE_THRESHOLD and not self.preserve_formulas

        if use_readonly:
            logger.info("Large file detected (%.1f MB) -- using read-only streaming mode",
                        file_size / (1024 * 1024))

        wb = load_workbook(
            str(input_path),
            data_only=not self.preserve_formulas,
            read_only=use_readonly,
        )

        sections: List[str] = []

        if self.include_metadata:
            sections.append(self._file_metadata(input_path, wb))

        sheets_to_process = self._get_sheets(wb)

        for sheet_name in sheets_to_process:
            sheet = wb[sheet_name]
            if sheet.sheet_state == "hidden" and not self.include_hidden_sheets:
                continue
            sections.append(self._convert_sheet(sheet, sheet_name, use_readonly))

        wb.close()

        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text("\n\n".join(sections), encoding="utf-8")

        logger.info("Converted %s (%d sheets)", input_path.name, len(sheets_to_process))
        return ConversionResult(
            input_path=input_path,
            output_path=output_path,
            success=True,
            metadata={"sheets": len(sheets_to_process)},
        )

    # -- internal helpers -----------------------------------------------------

    def _file_metadata(self, input_path: Path, wb: Any) -> str:
        lines = [
            "# Excel to Markdown Conversion",
            f"\n**Source File:** `{input_path.name}`",
            f"**Converted:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            f"**Total Sheets:** {len(wb.sheetnames)}",
        ]
        if wb.properties:
            if wb.properties.creator:
                lines.append(f"**Created By:** {wb.properties.creator}")
            if wb.properties.created:
                lines.append(f"**Creation Date:** {wb.properties.created}")
            if wb.properties.modified:
                lines.append(f"**Last Modified:** {wb.properties.modified}")
        lines.append("\n---")
        return "\n".join(lines)

    def _get_sheets(self, wb: Any) -> List[str]:
        if self.sheets:
            available = set(wb.sheetnames)
            missing = set(self.sheets) - available
            if missing:
                logger.warning("Sheets not found: %s", missing)
            return [s for s in self.sheets if s in available]
        return list(wb.sheetnames)

    def _convert_sheet(self, sheet: Any, name: str, readonly: bool) -> str:
        from tabulate import tabulate as _tabulate

        content = [f"## Sheet: {name}"]

        if self.include_metadata:
            meta = self._sheet_metadata(sheet, readonly)
            if meta:
                content.append(meta)

        # Read data
        table_data, data_range = self._read_data(sheet, readonly)
        if not table_data:
            content.append("\n*Empty sheet*")
            return "\n".join(content)

        # Merged cells info (not available in read-only mode)
        if not readonly and data_range:
            merged = self._merged_cells_info(sheet, data_range)
            if merged:
                content.append("\n**Note:** This sheet contains merged cells:")
                for info in merged:
                    content.append(f"- {info}")

        # Build Markdown table
        headers = table_data[0]
        looks_like_headers = any(
            h and not h.replace(".", "").replace("-", "").replace(" ", "").isdigit()
            for h in headers
        )

        if looks_like_headers and len(table_data) > 1:
            md_table = _tabulate(table_data[1:], headers=headers, tablefmt="pipe")
        else:
            from openpyxl.utils import get_column_letter
            col_names = [f"Column {get_column_letter(i + 1)}" for i in range(len(headers))]
            md_table = _tabulate(table_data, headers=col_names, tablefmt="pipe")

        content.append(f"\n{md_table}")

        # Hyperlinks (not available in read-only mode)
        if not readonly and data_range:
            links = self._extract_hyperlinks(sheet, data_range)
            if links:
                content.append("\n### Links")
                for ref, url in links.items():
                    content.append(f"- Cell {ref}: [{url}]({url})")

        # Formulas
        if self.preserve_formulas and not readonly and data_range:
            formulas = self._extract_formulas(sheet, data_range)
            if formulas:
                content.append("\n### Formulas")
                for ref, formula in formulas.items():
                    content.append(f"- {ref}: `{formula}`")

        return "\n".join(content)

    def _sheet_metadata(self, sheet: Any, readonly: bool) -> str:
        parts: List[str] = []
        if not readonly and sheet.max_row and sheet.max_column:
            parts.append(f"*Dimensions: {sheet.max_row} rows x {sheet.max_column} columns*")
        if not readonly and hasattr(sheet, "protection") and sheet.protection.sheet:
            parts.append("*Protected: Yes*")
        if sheet.sheet_state == "hidden":
            parts.append("*Hidden sheet*")
        return " | ".join(parts)

    def _read_data(
        self, sheet: Any, readonly: bool
    ) -> Tuple[List[List[str]], Optional[Tuple[int, int, int, int]]]:
        """Read cell values.  Returns (rows, data_range) or ([], None)."""
        if readonly:
            return self._read_data_streaming(sheet)
        return self._read_data_random(sheet)

    def _read_data_random(
        self, sheet: Any
    ) -> Tuple[List[List[str]], Optional[Tuple[int, int, int, int]]]:
        """Random-access read (standard mode)."""
        if not sheet.max_row or not sheet.max_column:
            return [], None

        min_row, max_row = sheet.max_row, 1
        min_col, max_col = sheet.max_column, 1

        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    min_row = min(min_row, cell.row)
                    max_row = max(max_row, cell.row)
                    min_col = min(min_col, cell.column)
                    max_col = max(max_col, cell.column)

        if max_row < min_row:
            return [], None

        data_range = (min_row, min_col, max_row, max_col)
        table: List[List[str]] = []
        for r in range(min_row, max_row + 1):
            row_data: List[str] = []
            for c in range(min_col, max_col + 1):
                cell = sheet.cell(row=r, column=c)
                val = self._cell_value(cell)
                if self.max_column_width and len(val) > self.max_column_width:
                    val = val[: self.max_column_width - 3] + "..."
                row_data.append(val)
            table.append(row_data)

        return table, data_range

    def _read_data_streaming(self, sheet: Any) -> Tuple[List[List[str]], None]:
        """Streaming read (read-only mode) for large workbooks."""
        table: List[List[str]] = []
        for row in sheet.iter_rows():
            row_data: List[str] = []
            for cell in row:
                val = self._cell_value(cell)
                if self.max_column_width and len(val) > self.max_column_width:
                    val = val[: self.max_column_width - 3] + "..."
                row_data.append(val)
            if any(v != self.empty_cell_placeholder for v in row_data):
                table.append(row_data)
        return table, None

    def _cell_value(self, cell: Any) -> str:
        """Format a single cell value as a string."""
        from openpyxl.cell.cell import MergedCell

        if isinstance(cell, MergedCell):
            return self.empty_cell_placeholder
        if cell.value is None:
            return self.empty_cell_placeholder

        value = cell.value

        if hasattr(value, "strftime"):
            return value.strftime("%Y-%m-%d")

        if isinstance(value, (int, float)):
            fmt = getattr(cell, "number_format", None) or ""
            if "%" in fmt:
                return f"{value:.2%}"
            if isinstance(value, float):
                return str(int(value)) if value == int(value) else f"{value:.2f}"
            return str(value)

        if isinstance(value, bool):
            return "TRUE" if value else "FALSE"

        s = str(value).replace("|", "\\|")
        return " ".join(s.split())

    def _merged_cells_info(
        self, sheet: Any, data_range: Tuple[int, int, int, int]
    ) -> List[str]:
        from openpyxl.utils import get_column_letter

        min_row, min_col, max_row, max_col = data_range
        info: List[str] = []
        for mr in sheet.merged_cells.ranges:
            if (
                mr.min_row <= max_row
                and mr.max_row >= min_row
                and mr.min_col <= max_col
                and mr.max_col >= min_col
            ):
                start = f"{get_column_letter(mr.min_col)}{mr.min_row}"
                end = f"{get_column_letter(mr.max_col)}{mr.max_row}"
                info.append(f"{start}:{end}")
        return info

    def _extract_hyperlinks(
        self, sheet: Any, data_range: Tuple[int, int, int, int]
    ) -> Dict[str, str]:
        from openpyxl.utils import get_column_letter

        min_row, min_col, max_row, max_col = data_range
        links: Dict[str, str] = {}
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = sheet.cell(row=r, column=c)
                if cell.hyperlink:
                    ref = f"{get_column_letter(c)}{r}"
                    links[ref] = cell.hyperlink.target
        return links

    def _extract_formulas(
        self, sheet: Any, data_range: Tuple[int, int, int, int]
    ) -> Dict[str, str]:
        from openpyxl.utils import get_column_letter

        min_row, min_col, max_row, max_col = data_range
        formulas: Dict[str, str] = {}
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                cell = sheet.cell(row=r, column=c)
                if cell.value and isinstance(cell.value, str) and cell.value.startswith("="):
                    ref = f"{get_column_letter(c)}{r}"
                    formulas[ref] = cell.value
        return formulas

    # -- CLI ------------------------------------------------------------------

    @classmethod
    def add_arguments(cls, parser) -> None:
        parser.add_argument("--sheets", nargs="+", help="Specific sheet names to convert")
        parser.add_argument("--preserve-formulas", action="store_true",
                            help="Show formulas instead of calculated values")
        parser.add_argument("--no-metadata", action="store_true",
                            help="Exclude file and sheet metadata")
        parser.add_argument("--include-hidden", action="store_true",
                            help="Include hidden sheets")
        parser.add_argument("--max-width", type=int, default=0,
                            help="Maximum column width (0 = no limit, default: 0)")

    @classmethod
    def from_args(cls, args) -> "XlsxConverter":
        return cls(
            sheets=getattr(args, "sheets", None),
            preserve_formulas=getattr(args, "preserve_formulas", False),
            include_metadata=not getattr(args, "no_metadata", False),
            include_hidden_sheets=getattr(args, "include_hidden", False),
            max_column_width=getattr(args, "max_width", 0),
        )
